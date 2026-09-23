from pathlib import Path

from openpyxl import load_workbook

from vertere_api.importacao.domain import (
    CodigoInconsistencia,
    DadosPlanilha,
    InconsistenciaImportacao,
    LinhaAtendimentoPlanilha,
    LinhaClinicaPlanilha,
    LinhaVeterinarioPlanilha,
    SeveridadeInconsistencia,
)


_CABECALHOS = {
    "Cadastro Clínicas": (
        2,
        ("Clínica", "CNPJ", "Endereço", "Telefone", "E-mail", "Status"),
    ),
    "Cadastro Veterinários": (
        2,
        ("Clínica", "Nome", "CRMV", "Telefone", "E-mail (laudos)", "Status"),
    ),
    "Dados": (
        4,
        (
            "Clínica", "Data", "Hora", "Ano", "Nº", "Protocolo", "Método de Coleta",
            "Veterinário", "Paciente", "Espécie", "Raça", "Sexo", "Idade", "Proprietário",
            "Tipo de Exame", "Exame", "Tipo", "Valor", "Desconto", "Adicional", "Valor Total",
        ),
    ),
}


class EstruturaPlanilhaInvalida(Exception):
    def __init__(self, inconsistencias: tuple[InconsistenciaImportacao, ...]) -> None:
        self.inconsistencias = inconsistencias
        super().__init__("A planilha não possui a estrutura esperada")


def _erro_estrutura(
    codigo: CodigoInconsistencia, aba: str, linha: int | None, mensagem: str
) -> InconsistenciaImportacao:
    return InconsistenciaImportacao(
        severidade=SeveridadeInconsistencia.ERRO,
        codigo=codigo,
        aba=aba,
        linha=linha,
        coluna=None,
        mensagem=mensagem,
    )


def _mapear_linha(cabecalhos: tuple[str, ...], valores: tuple[object, ...]) -> dict[str, object]:
    return {cabecalho: valores[indice] for indice, cabecalho in enumerate(cabecalhos)}


def ler_planilha(caminho: str | Path) -> DadosPlanilha:
    livro = load_workbook(caminho, read_only=True, data_only=True)
    inconsistencias: list[InconsistenciaImportacao] = []
    indices: dict[str, tuple[int, tuple[str, ...]]] = {}

    for aba, (linha_cabecalho, esperados) in _CABECALHOS.items():
        if aba not in livro.sheetnames:
            inconsistencias.append(
                _erro_estrutura(
                    CodigoInconsistencia.ABA_AUSENTE,
                    aba,
                    None,
                    "Aba obrigatória ausente",
                )
            )
            continue
        planilha = livro[aba]
        celulas_cabecalho = next(
            planilha.iter_rows(
                min_row=linha_cabecalho,
                max_row=linha_cabecalho,
                max_col=len(esperados),
            ),
            (),
        )
        encontrados = tuple(
            "" if celula.value is None else str(celula.value).strip()
            for celula in celulas_cabecalho
        )
        if encontrados != esperados:
            inconsistencias.append(
                _erro_estrutura(
                    CodigoInconsistencia.CABECALHO_AUSENTE,
                    aba,
                    linha_cabecalho,
                    "Cabeçalhos obrigatórios ausentes ou fora de ordem",
                )
            )
            continue
        indices[aba] = (linha_cabecalho, esperados)

    if inconsistencias:
        livro.close()
        raise EstruturaPlanilhaInvalida(tuple(inconsistencias))

    clinicas: list[LinhaClinicaPlanilha] = []
    planilha_clinicas = livro["Cadastro Clínicas"]
    inicio, cabecalhos = indices["Cadastro Clínicas"]
    for numero, valores in enumerate(
        planilha_clinicas.iter_rows(
            min_row=inicio + 1, max_col=len(cabecalhos), values_only=True
        ),
        start=inicio + 1,
    ):
        if not any(valor is not None for valor in valores):
            continue
        item = _mapear_linha(cabecalhos, valores)
        clinicas.append(
            LinhaClinicaPlanilha(
                linha=numero,
                nome=item["Clínica"],
                cnpj=item["CNPJ"],
                endereco=item["Endereço"],
                telefone=item["Telefone"],
                email=item["E-mail"],
                status=item["Status"],
            )
        )

    veterinarios: list[LinhaVeterinarioPlanilha] = []
    planilha_veterinarios = livro["Cadastro Veterinários"]
    inicio, cabecalhos = indices["Cadastro Veterinários"]
    for numero, valores in enumerate(
        planilha_veterinarios.iter_rows(
            min_row=inicio + 1, max_col=len(cabecalhos), values_only=True
        ),
        start=inicio + 1,
    ):
        if not any(valor is not None for valor in valores):
            continue
        item = _mapear_linha(cabecalhos, valores)
        veterinarios.append(
            LinhaVeterinarioPlanilha(
                linha=numero,
                clinica=item["Clínica"],
                nome=item["Nome"],
                crmv=item["CRMV"],
                telefone=item["Telefone"],
                email=item["E-mail (laudos)"],
                status=item["Status"],
            )
        )

    atendimentos: list[LinhaAtendimentoPlanilha] = []
    planilha_dados = livro["Dados"]
    inicio, cabecalhos = indices["Dados"]
    for numero, valores in enumerate(
        planilha_dados.iter_rows(
            min_row=inicio + 1, max_col=len(cabecalhos), values_only=True
        ),
        start=inicio + 1,
    ):
        if not any(valor is not None for valor in valores):
            continue
        item = _mapear_linha(cabecalhos, valores)
        atendimentos.append(
            LinhaAtendimentoPlanilha(
                linha=numero,
                clinica=item["Clínica"],
                data=item["Data"],
                hora=item["Hora"],
                numero=item["Nº"],
                protocolo=item["Protocolo"],
                metodo_coleta=item["Método de Coleta"],
                veterinario=item["Veterinário"],
                paciente=item["Paciente"],
                especie=item["Espécie"],
                raca=item["Raça"],
                sexo=item["Sexo"],
                idade=item["Idade"],
                proprietario=item["Proprietário"],
                categoria_exame=item["Tipo de Exame"],
                exame=item["Exame"],
                tipo_atendimento=item["Tipo"],
                valor=item["Valor"],
                desconto=item["Desconto"],
                adicional=item["Adicional"],
                valor_total=item["Valor Total"],
            )
        )

    livro.close()
    return DadosPlanilha(
        clinicas=tuple(clinicas),
        veterinarios=tuple(veterinarios),
        atendimentos=tuple(atendimentos),
    )
