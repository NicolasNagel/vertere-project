from datetime import time, timedelta
from pathlib import Path

import pytest
from openpyxl import Workbook

from vertere_api.importacao.xlsx import EstruturaPlanilhaInvalida, ler_planilha


CABECALHOS_CLINICAS = [
    "Clínica", "CNPJ", "Endereço", "Telefone", "E-mail", "Status", "Observação"
]
CABECALHOS_VETERINARIOS = [
    "Clínica", "Nome", "CRMV", "Telefone", "E-mail (laudos)", "Status", "Observação"
]
CABECALHOS_DADOS = [
    "Clínica", "Data", "Hora", "Ano", "Nº", "Protocolo", "Método de Coleta",
    "Veterinário", "Paciente", "Espécie", "Raça", "Sexo", "Idade", "Proprietário",
    "Tipo de Exame", "Exame", "Tipo", "Valor", "Desconto", "Adicional", "Valor Total",
]


def _salvar_planilha_valida(caminho: Path, *, hora: object = "10:30") -> None:
    workbook = Workbook()
    clinicas = workbook.active
    clinicas.title = "Cadastro Clínicas"
    clinicas.append(["CADASTRO DE CLÍNICAS"])
    clinicas.append(CABECALHOS_CLINICAS)
    clinicas.append(["Clínica A", "12345678000190", "Rua A", "47999990000", "a@test", "Ativo", None])

    veterinarios = workbook.create_sheet("Cadastro Veterinários")
    veterinarios.append(["CADASTRO DE VETERINÁRIOS"])
    veterinarios.append(CABECALHOS_VETERINARIOS)
    veterinarios.append(["Clínica A", "Dra. Ana", "12345", None, None, "Ativo", None])

    dados = workbook.create_sheet("Dados")
    dados.append(["REGISTRO DE ATENDIMENTOS"])
    dados.append([])
    dados.append([])
    dados.append(CABECALHOS_DADOS)
    dados.append([
        "Clínica A", "10/01/2026", hora, 2026, 1, "PROTO-1", "Motoboy",
        "Dra. Ana", "Tobias", "Canino", "SRD", "Macho", 4, "José",
        "Hematológicos", "Hemograma", "Normal", 50, None, None, 50,
    ])
    workbook.save(caminho)


def test_le_planilha_minima_com_as_tres_abas(tmp_path: Path) -> None:
    caminho = tmp_path / "historico.xlsx"
    _salvar_planilha_valida(caminho)

    dados = ler_planilha(caminho)

    assert len(dados.clinicas) == 1
    assert dados.clinicas[0].linha == 3
    assert dados.clinicas[0].nome == "Clínica A"
    assert len(dados.veterinarios) == 1
    assert dados.veterinarios[0].linha == 3
    assert len(dados.atendimentos) == 1
    assert dados.atendimentos[0].linha == 5
    assert dados.atendimentos[0].protocolo == "PROTO-1"


def test_normaliza_hora_excel_representada_como_duracao(tmp_path: Path) -> None:
    caminho = tmp_path / "hora-duracao.xlsx"
    _salvar_planilha_valida(caminho, hora=timedelta(hours=17, minutes=11))

    dados = ler_planilha(caminho)

    assert dados.atendimentos[0].hora == time(17, 11)


def test_rejeita_aba_ou_cabecalho_ausente(tmp_path: Path) -> None:
    caminho = tmp_path / "invalida.xlsx"
    workbook = Workbook()
    workbook.active.title = "Cadastro Clínicas"
    workbook.active.append(["cabeçalho errado"])
    workbook.save(caminho)

    with pytest.raises(EstruturaPlanilhaInvalida) as captura:
        ler_planilha(caminho)

    codigos = {item.codigo.value for item in captura.value.inconsistencias}
    assert "aba_ausente" in codigos
    assert "cabecalho_ausente" in codigos
