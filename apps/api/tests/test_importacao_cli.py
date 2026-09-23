import json
from pathlib import Path
from uuid import uuid5

from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from vertere_api.atendimentos.models import AtendimentoModel
from vertere_api.clinicas.models import ClinicaModel
from vertere_api.db import Base
from vertere_api.importacao.cli import executar
from vertere_api.importacao.domain import NAMESPACE_CLINICAS


def _planilha(caminho: Path, *, cnpj: str | None = "12345678000190") -> None:
    workbook = Workbook()
    clinicas = workbook.active
    clinicas.title = "Cadastro Clínicas"
    clinicas.append(["CADASTRO"])
    clinicas.append(["Clínica", "CNPJ", "Endereço", "Telefone", "E-mail", "Status"])
    clinicas.append(["Clínica A", cnpj, "Rua A", "", "", "Ativo"])

    veterinarios = workbook.create_sheet("Cadastro Veterinários")
    veterinarios.append(["CADASTRO"])
    veterinarios.append(["Clínica", "Nome", "CRMV", "Telefone", "E-mail (laudos)", "Status"])
    veterinarios.append(["Clínica A", "Dra. Ana", "12345", "", "", "Ativo"])

    dados = workbook.create_sheet("Dados")
    dados.append(["REGISTRO"])
    dados.append([])
    dados.append([])
    dados.append([
        "Clínica", "Data", "Hora", "Ano", "Nº", "Protocolo", "Método de Coleta",
        "Veterinário", "Paciente", "Espécie", "Raça", "Sexo", "Idade", "Proprietário",
        "Tipo de Exame", "Exame", "Tipo", "Valor", "Desconto", "Adicional", "Valor Total",
    ])
    dados.append([
        "Clínica A", "10/01/2026", "10:30", 2026, 1, "PROTO-1", "Motoboy",
        "Dra. Ana", "Tobias", "Canino", "SRD", "Macho", 4, "José",
        "Hematológicos", "Hemograma", "Normal", 50, None, None, 50,
    ])
    workbook.save(caminho)


def test_dry_run_gera_relatorio_sem_exigir_banco(tmp_path: Path, capsys) -> None:
    planilha = tmp_path / "historico.xlsx"
    relatorio = tmp_path / "relatorio.json"
    _planilha(planilha)

    codigo = executar([str(planilha), "--relatorio", str(relatorio)])

    conteudo = json.loads(relatorio.read_text(encoding="utf-8"))
    assert codigo == 0
    assert conteudo["modo"] == "dry-run"
    assert conteudo["aplicavel"] is True
    assert conteudo["contadores"]["atendimentos"] == 1
    assert conteudo["totais_inconsistencias"] == {"erros": 0, "avisos": 0}
    assert json.loads(capsys.readouterr().out) == conteudo
    assert "Tobias" not in relatorio.read_text(encoding="utf-8")


def test_inconsistencia_bloqueia_aplicacao_e_e_reportada(tmp_path: Path) -> None:
    planilha = tmp_path / "invalida.xlsx"
    relatorio = tmp_path / "relatorio.json"
    _planilha(planilha, cnpj=None)

    codigo = executar([str(planilha), "--relatorio", str(relatorio)])

    conteudo = json.loads(relatorio.read_text(encoding="utf-8"))
    assert codigo == 1
    assert conteudo["aplicavel"] is False
    assert conteudo["totais_inconsistencias"]["erros"] > 0
    assert any(item["codigo"] == "cnpj_invalido" for item in conteudo["inconsistencias"])


def test_aplicar_persiste_plano_valido(tmp_path: Path) -> None:
    planilha = tmp_path / "historico.xlsx"
    relatorio = tmp_path / "relatorio.json"
    banco = tmp_path / "destino.sqlite"
    _planilha(planilha)
    url = f"sqlite:///{banco.as_posix()}"
    engine = create_engine(url)
    Base.metadata.create_all(engine)

    codigo = executar(
        [
            str(planilha),
            "--relatorio", str(relatorio),
            "--aplicar",
            "--banco", url,
        ]
    )

    with Session(engine) as session:
        quantidade = session.scalar(select(func.count()).select_from(AtendimentoModel))
    engine.dispose()
    assert codigo == 0
    assert quantidade == 1
    assert json.loads(relatorio.read_text(encoding="utf-8"))["modo"] == "aplicado"


def test_colisao_ao_aplicar_faz_rollback_e_gera_relatorio(tmp_path: Path) -> None:
    planilha = tmp_path / "historico.xlsx"
    relatorio = tmp_path / "relatorio.json"
    banco = tmp_path / "destino.sqlite"
    _planilha(planilha)
    url = f"sqlite:///{banco.as_posix()}"
    motor = create_engine(url)
    Base.metadata.create_all(motor)
    with Session(motor) as sessao:
        sessao.add(
            ClinicaModel(
                id=str(uuid5(NAMESPACE_CLINICAS, "12345678000190")),
                nome="Conflito",
                cnpj="99999999000199",
                endereco="",
                telefone="",
                email="",
                ativo=True,
            )
        )
        sessao.commit()
    motor.dispose()

    codigo = executar(
        [str(planilha), "--relatorio", str(relatorio), "--aplicar", "--banco", url]
    )

    conteudo = json.loads(relatorio.read_text(encoding="utf-8"))
    assert codigo == 1
    assert conteudo["modo"] == "falhou"
    assert conteudo["aplicavel"] is False
    assert conteudo["totais_inconsistencias"]["erros"] == 1
    assert conteudo["inconsistencias"][0]["codigo"] == "colisao_destino"
